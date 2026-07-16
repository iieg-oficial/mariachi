import csv
import io
import zipfile

from openpyxl import load_workbook

from app.services.sieej.xlsx_service import (
    build_envios_csv,
    build_envios_tables,
    build_envios_xlsx,
)

DEF_VIGENTE = {
    "version": 1,
    "steps": [
        {
            "id": "general",
            "type": "form",
            "title": "General",
            "fields": [
                {"name": "razon", "label": "Razon social", "type": "text"},
            ],
        }
    ],
}

DEF_HISTORICA = {
    "version": 1,
    "steps": [
        {
            "id": "general",
            "type": "form",
            "title": "General",
            "fields": [
                {"name": "razon", "label": "Razon social", "type": "text"},
                {"name": "viejo", "label": "Campo viejo", "type": "text"},
            ],
        }
    ],
}

DEF_CON_REPEATER = {
    "version": 1,
    "steps": [
        {
            "id": "general",
            "type": "form",
            "title": "General",
            "fields": [{"name": "razon", "label": "Razon social", "type": "text"}],
        },
        {
            "id": "bases",
            "type": "repeater",
            "title": "Bases de datos",
            "fields": [{"name": "nombre_bd", "label": "Nombre", "type": "text"}],
        },
    ],
}


def _envio(**overrides):
    base = {
        "id": 1,
        "usuario_nombre": "Ana",
        "usuario_email": "ana@test.com",
        "estado": "enviado",
        "formulario_version": 2,
        "enviado_en": "2026-07-16 10:00",
        "datos": {"general": {"razon": "ACME", "viejo": "dato huerfano"}},
        "definicion": DEF_VIGENTE,
    }
    base.update(overrides)
    return base


def test_tables_incluyen_campos_eliminados_via_historicas():
    tables = build_envios_tables(
        [_envio()],
        definiciones_historicas=[DEF_HISTORICA],
        definicion_vigente=DEF_VIGENTE,
    )
    headers = tables[0]["headers"]
    assert "Campo viejo (eliminado)" in headers
    assert "Versión" in headers
    row = tables[0]["rows"][0]
    assert row[headers.index("Campo viejo (eliminado)")] == "dato huerfano"
    assert row[headers.index("Versión")] == 2
    assert row[headers.index("Razon social")] == "ACME"


def test_tables_sin_historicas_no_marcan_eliminados():
    tables = build_envios_tables([_envio()], definicion_vigente=DEF_VIGENTE)
    headers = tables[0]["headers"]
    assert "Campo viejo (eliminado)" not in headers
    assert "Razon social" in headers


def test_xlsx_incluye_columna_de_campo_eliminado():
    contenido = build_envios_xlsx(
        [_envio()],
        definiciones_historicas=[DEF_HISTORICA],
        definicion_vigente=DEF_VIGENTE,
    )
    wb = load_workbook(io.BytesIO(contenido))
    ws = wb["Envios"]
    headers = [c.value for c in ws[1]]
    assert "Campo viejo (eliminado)" in headers
    assert ws[2][headers.index("Campo viejo (eliminado)")].value == "dato huerfano"


def test_csv_plano_sin_repeater():
    contenido, es_zip = build_envios_csv([_envio()], definicion_vigente=DEF_VIGENTE)
    assert es_zip is False
    filas = list(csv.reader(io.StringIO(contenido.decode("utf-8-sig"))))
    assert "Razon social" in filas[0]
    assert filas[1][filas[0].index("Razon social")] == "ACME"


def test_csv_zip_con_repeater():
    envio = _envio(
        datos={
            "general": {"razon": "ACME"},
            "bases": [{"nombre_bd": "BD1"}, {"nombre_bd": "BD2"}],
        },
        definicion=DEF_CON_REPEATER,
    )
    contenido, es_zip = build_envios_csv([envio], definicion_vigente=DEF_CON_REPEATER)
    assert es_zip is True
    with zipfile.ZipFile(io.BytesIO(contenido)) as zf:
        nombres = sorted(zf.namelist())
        assert nombres == ["bases_de_datos.csv", "envios.csv"]
        filas = list(
            csv.reader(io.StringIO(zf.read("bases_de_datos.csv").decode("utf-8-sig")))
        )
        assert filas[0][3] == "Nombre"
        assert [f[3] for f in filas[1:]] == ["BD1", "BD2"]
