from __future__ import annotations

import csv
import io
import logging
from datetime import datetime
from typing import Any

from sqlalchemy import text
from sqlalchemy.engine import Connection

from app.services.grid_batch import GridSpec

logger = logging.getLogger(__name__)

HISTORY_HEADERS = (
    'Capa',
    'Campo',
    'Valor anterior',
    'Valor nuevo',
    'Responsable',
    'Fecha',
    'Origen',
)

SOURCE_LABELS = {
    'grid': 'Tabla de captura',
    'formulario': 'Ficha de la capa',
    'ingesta': 'Ingesta masiva',
    'backfill': 'Registro inicial',
}


def _format_datetime(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M')
    return '' if value is None else str(value)


def _cell(value: Any) -> Any:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Sí' if value else 'No'
    return value


def build_rows_sheet(spec: GridSpec, rows: list[dict]) -> tuple[list[str], list[list]]:
    columns = [meta for meta in spec.columns_meta]
    headers = [meta['title'] for meta in columns] + ['Responsable', 'Última modificación']
    body = [
        [_cell(row.get(meta['key'])) for meta in columns]
        + [_cell(row.get('updated_by')), _format_datetime(row.get('updated_at'))]
        for row in rows
    ]
    return headers, body


def fetch_history(
    conn: Connection,
    spec: GridSpec,
    row_keys: list[str] | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    limit: int = 50000,
) -> list[dict]:
    if not spec.history_table:
        return []

    clauses = ['resource = :resource']
    params: dict[str, Any] = {'resource': spec.key, 'limit': limit}
    if row_keys:
        clauses.append('row_key = ANY(:row_keys)')
        params['row_keys'] = row_keys
    if date_from:
        clauses.append('changed_at >= :date_from')
        params['date_from'] = date_from
    if date_to:
        clauses.append('changed_at < (CAST(:date_to AS date) + 1)')
        params['date_to'] = date_to

    if conn.execute(text('SELECT to_regclass(:table)'), {'table': spec.history_table}).scalar() is None:
        logger.warning(
            'grid_export: %s no existe todavia; se exporta sin historial', spec.history_table
        )
        return []

    result = conn.execute(
        text(
            'SELECT row_key, column_key, from_value, to_value, changed_by, changed_at, source '
            f'FROM {spec.history_table} '
            f'WHERE {" AND ".join(clauses)} '
            'ORDER BY changed_at DESC, id DESC LIMIT :limit'
        ),
        params,
    ).mappings().all()
    return [dict(row) for row in result]


def build_history_sheet(spec: GridSpec, history: list[dict]) -> tuple[list[str], list[list]]:
    titles = {meta['key']: meta['title'] for meta in spec.columns_meta}
    body = [
        [
            entry['row_key'],
            titles.get(entry['column_key'], entry['column_key']),
            _cell(entry['from_value']),
            _cell(entry['to_value']),
            _cell(entry['changed_by']),
            _format_datetime(entry['changed_at']),
            SOURCE_LABELS.get(entry['source'], entry['source']),
        ]
        for entry in history
    ]
    return list(HISTORY_HEADERS), body


def to_csv(headers: list[str], body: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(body)
    return buffer.getvalue().encode('utf-8-sig')


def to_xlsx(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError('openpyxl no instalado') from exc

    workbook = Workbook()
    workbook.remove(workbook.active)

    for title, headers, body in sheets:
        sheet = workbook.create_sheet(title=title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in body:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        for index, header in enumerate(headers, start=1):
            width = max(12, min(52, len(str(header)) + 6))
            sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
