"""Exporta los envios de un formulario a Excel (openpyxl) o CSV.

Tabla principal `Envios`: una fila por envio, columnas de metadatos +
un campo por cada campo de los pasos tipo `form`. Cada paso `repeater`
genera una tabla aparte con una fila por item. En Excel cada tabla es
una hoja; en CSV cada tabla es un archivo (ZIP cuando hay mas de una).

Las columnas se arman con la union de los snapshots de los envios, la
definicion vigente y las definiciones historicas archivadas en
`formulario_version`. Asi los valores de campos que ya no existen en la
definicion vigente siguen saliendo (permanecen en `datos` aunque nada
los renderice); sus encabezados llevan el sufijo "(eliminado)".
"""
from __future__ import annotations

import csv
import json
import re
import zipfile
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.sieej.definicion_validator import CLAVE_ETIQUETA
from app.services.sieej.export_format import format_value

ESTADO_LABEL = {
    "en_proceso": "En proceso",
    "enviado": "Enviado",
    "expirado": "Expirado",
}


def _cell(field: dict[str, Any], value: Any) -> str:
    return format_value(field, value) or ""


def _hist_valor(valor: Any) -> str:
    """Serializa un valor de historial (JSON arbitrario) para una celda."""
    if valor is None:
        return ""
    if isinstance(valor, (str, int, float, bool)):
        return str(valor)
    return json.dumps(valor, ensure_ascii=False)


_HISTORIAL_HEADERS = [
    "Envio ID",
    "Usuario",
    "Versión",
    "Campo",
    "Valor anterior",
    "Valor nuevo",
    "Actor",
    "Fecha",
]


def _historial_table(historial: list[dict[str, Any]]) -> dict[str, Any]:
    rows = [
        [
            h.get("envio_id"),
            h.get("usuario") or "",
            h.get("version") or "",
            h.get("campo") or "",
            _hist_valor(h.get("valor_anterior")),
            _hist_valor(h.get("valor_nuevo")),
            h.get("actor") or "",
            h.get("fecha") or "",
        ]
        for h in historial
    ]
    return {
        "title": "Historial de cambios",
        "headers": _HISTORIAL_HEADERS,
        "rows": rows,
    }


def _form_fields(definiciones: list[dict[str, Any]]) -> list[tuple[dict, dict]]:
    seen: dict[tuple, tuple[dict, dict]] = {}
    order: list[tuple] = []
    for definicion in definiciones:
        for step in (definicion or {}).get("steps", []):
            if step.get("type") in ("summary", "repeater"):
                continue
            for field in step.get("fields", []):
                if field.get("type") == "info":
                    continue
                key = (step.get("id"), field.get("name"))
                if key not in seen:
                    seen[key] = (step, field)
                    order.append(key)
    return [seen[k] for k in order]


def _repeater_steps(definiciones: list[dict[str, Any]]) -> list[tuple[dict, list[dict]]]:
    steps: dict[Any, dict] = {}
    order: list[Any] = []
    fields: dict[Any, tuple[dict, list]] = {}
    for definicion in definiciones:
        for step in (definicion or {}).get("steps", []):
            if step.get("type") != "repeater":
                continue
            sid = step.get("id")
            if sid not in steps:
                steps[sid] = step
                order.append(sid)
                fields[sid] = ({}, [])
            fmap, forder = fields[sid]
            for field in step.get("fields", []):
                if field.get("type") == "info":
                    continue
                fname = field.get("name")
                if fname not in fmap:
                    fmap[fname] = field
                    forder.append(fname)
    return [(steps[sid], [fields[sid][0][n] for n in fields[sid][1]]) for sid in order]


def _paths_vigentes(definicion: dict[str, Any] | None) -> set[tuple] | None:
    if not definicion:
        return None
    paths: set[tuple] = set()
    for step in definicion.get("steps", []) or []:
        for field in step.get("fields", []) or []:
            paths.add((step.get("id"), field.get("name")))
    return paths


def _header_label(
    field: dict[str, Any], step_id: Any, vigentes: set[tuple] | None
) -> str:
    label = field.get("label") or field.get("name")
    if vigentes is not None and (step_id, field.get("name")) not in vigentes:
        return f"{label} (eliminado)"
    return label


def build_envios_tables(
    envios: list[dict[str, Any]],
    *,
    definiciones_historicas: list[dict[str, Any]] | None = None,
    definicion_vigente: dict[str, Any] | None = None,
    historial: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Construye las tablas del export como `{title, headers, rows}`."""
    definiciones = [e.get("definicion") or {} for e in envios]
    if definicion_vigente:
        definiciones.append(definicion_vigente)
    definiciones.extend(definiciones_historicas or [])
    vigentes = _paths_vigentes(definicion_vigente)

    tables: list[dict[str, Any]] = []

    form_fields = _form_fields(definiciones)
    headers = ["ID", "Usuario", "Email", "Estado", "Versión", "Enviado"] + [
        _header_label(f, s.get("id"), vigentes) for s, f in form_fields
    ]
    rows: list[list[Any]] = []
    for envio in envios:
        datos = envio.get("datos") or {}
        row: list[Any] = [
            envio.get("id"),
            envio.get("usuario_nombre") or "",
            envio.get("usuario_email") or "",
            ESTADO_LABEL.get(envio.get("estado"), envio.get("estado") or ""),
            envio.get("formulario_version") or "",
            envio.get("enviado_en") or "",
        ]
        for step, field in form_fields:
            raw = datos.get(step.get("id"))
            step_data = raw if isinstance(raw, dict) else {}
            row.append(_cell(field, step_data.get(field.get("name"))))
        rows.append(row)
    tables.append({"title": "Envios", "headers": headers, "rows": rows})

    for step, fields in _repeater_steps(definiciones):
        sid = step.get("id")
        rep_headers = ["Envio ID", "Usuario", "#"] + [
            _header_label(f, sid, vigentes) for f in fields
        ] + ["Nombre de la pestaña"]
        rep_rows: list[list[Any]] = []
        for envio in envios:
            items = (envio.get("datos") or {}).get(sid)
            if not isinstance(items, list):
                continue
            for i, item in enumerate(items):
                item = item if isinstance(item, dict) else {}
                rep_rows.append(
                    [envio.get("id"), envio.get("usuario_nombre") or "", i + 1]
                    + [_cell(f, item.get(f.get("name"))) for f in fields]
                    + [item.get(CLAVE_ETIQUETA) or ""]
                )
        tables.append(
            {"title": step.get("title") or sid, "headers": rep_headers, "rows": rep_rows}
        )

    if historial:
        tables.append(_historial_table(historial))

    return tables


def _unique_sheet_title(wb: Workbook, base: str) -> str:
    title = (base or "Datos")[:31]
    candidate = title
    i = 1
    while candidate in wb.sheetnames:
        suffix = f" ({i})"
        candidate = title[: 31 - len(suffix)] + suffix
        i += 1
    return candidate


def _bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


def build_envios_xlsx(
    envios: list[dict[str, Any]],
    *,
    definiciones_historicas: list[dict[str, Any]] | None = None,
    definicion_vigente: dict[str, Any] | None = None,
    historial: list[dict[str, Any]] | None = None,
) -> bytes:
    tables = build_envios_tables(
        envios,
        definiciones_historicas=definiciones_historicas,
        definicion_vigente=definicion_vigente,
        historial=historial,
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Envios"
    principal, *repeaters = tables

    ws.append(principal["headers"])
    _bold_header(ws)
    for row in principal["rows"]:
        ws.append(row)

    for table in repeaters:
        rs = wb.create_sheet(_unique_sheet_title(wb, table["title"]))
        rs.append(table["headers"])
        _bold_header(rs)
        for row in table["rows"]:
            rs.append(row)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _csv_bytes(headers: list[Any], rows: list[list[Any]]) -> bytes:
    sio = StringIO()
    writer = csv.writer(sio)
    writer.writerow(headers)
    writer.writerows(rows)
    return sio.getvalue().encode("utf-8-sig")


def _csv_filename(base: str, usados: set[str]) -> str:
    nombre = re.sub(r"[^\w-]+", "_", (base or "datos").lower()).strip("_") or "datos"
    candidato = nombre
    i = 1
    while candidato in usados:
        i += 1
        candidato = f"{nombre}_{i}"
    usados.add(candidato)
    return candidato


def build_envios_csv(
    envios: list[dict[str, Any]],
    *,
    definiciones_historicas: list[dict[str, Any]] | None = None,
    definicion_vigente: dict[str, Any] | None = None,
    historial: list[dict[str, Any]] | None = None,
) -> tuple[bytes, bool]:
    """Devuelve `(contenido, es_zip)`: CSV plano con una sola tabla, ZIP con
    un CSV por tabla cuando el formulario tiene pasos repeater o historial."""
    tables = build_envios_tables(
        envios,
        definiciones_historicas=definiciones_historicas,
        definicion_vigente=definicion_vigente,
        historial=historial,
    )
    if len(tables) == 1:
        return _csv_bytes(tables[0]["headers"], tables[0]["rows"]), False

    bio = BytesIO()
    usados: set[str] = set()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        for table in tables:
            nombre = _csv_filename(table["title"], usados)
            zf.writestr(f"{nombre}.csv", _csv_bytes(table["headers"], table["rows"]))
    return bio.getvalue(), True
