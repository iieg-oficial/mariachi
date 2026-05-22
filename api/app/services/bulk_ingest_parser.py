from __future__ import annotations

import csv
import io
import re
from datetime import date
from typing import Any, Iterator

TRUE_VALUES = {'1', 'true', 't', 'yes', 'sí', 'si', 'y', 'sÍ', 'verdadero'}
FALSE_VALUES = {'0', 'false', 'f', 'no', 'n', 'falso'}

_MONTH_ES = {
    'ene': 1, 'enero': 1,
    'feb': 2, 'febrero': 2,
    'mar': 3, 'marzo': 3,
    'abr': 4, 'abril': 4,
    'may': 5, 'mayo': 5,
    'jun': 6, 'junio': 6,
    'jul': 7, 'julio': 7,
    'ago': 8, 'agosto': 8,
    'sep': 9, 'sept': 9, 'septiembre': 9,
    'oct': 10, 'octubre': 10,
    'nov': 11, 'noviembre': 11,
    'dic': 12, 'diciembre': 12,
}

_DATE_ISO_FULL = re.compile(r'^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T\s].*)?$')
_DATE_YYYY_MM = re.compile(r'^(\d{4})[-/](\d{1,2})$')
_DATE_YYYY = re.compile(r'^(\d{4})$')
_DATE_DMY_NUM = re.compile(r'^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})$')
_DATE_DMY_NAMED = re.compile(r'^(\d{1,2})[/\s\-]+([A-Za-zÁÉÍÓÚáéíóú]+)[/\s\-]+(\d{2,4})$')
_DATE_MY_NAMED = re.compile(r'^([A-Za-zÁÉÍÓÚáéíóú]+)[/\s\-]+(\d{4})$')

_NUM_PREFIX_SYM = re.compile(r'^([$])\s*')
_NUM_SUFFIX_SYM = re.compile(r'\s*([%])$')


def _valid_ymd(y: int, mo: int, d: int) -> bool:
    try:
        date(y, mo, d)
    except ValueError:
        return False
    return True


def _expand_two_digit_year(y: int) -> int:
    return y + (2000 if y < 70 else 1900) if y < 100 else y


def parse_date_iso(value: Any) -> str | None:
    """Convierte entradas de fecha al estandar IIEG (ISO 8601, YYYY-MM-DD).

    Reconoce: ISO completo, YYYY-MM, YYYY, DD/MM/YYYY (es-MX), DD-MMM-YYYY,
    'Mes YYYY'. Si la entrada no parsea, devuelve el string limpio sin tocar
    (preserva data ya en BD con formatos no estandar).
    """
    s = clean(value)
    if s is None:
        return None

    m = _DATE_ISO_FULL.match(s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if _valid_ymd(y, mo, d):
            return f'{y:04d}-{mo:02d}-{d:02d}'

    m = _DATE_YYYY_MM.match(s)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12:
            return f'{y:04d}-{mo:02d}-01'

    m = _DATE_YYYY.match(s)
    if m:
        return f'{int(m.group(1)):04d}-01-01'

    m = _DATE_DMY_NUM.match(s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), _expand_two_digit_year(int(m.group(3)))
        if _valid_ymd(y, mo, d):
            return f'{y:04d}-{mo:02d}-{d:02d}'

    m = _DATE_DMY_NAMED.match(s)
    if m:
        d, mname, y = int(m.group(1)), m.group(2).lower(), _expand_two_digit_year(int(m.group(3)))
        mo = _MONTH_ES.get(mname) or _MONTH_ES.get(mname[:3])
        if mo and _valid_ymd(y, mo, d):
            return f'{y:04d}-{mo:02d}-{d:02d}'

    m = _DATE_MY_NAMED.match(s)
    if m:
        mname, y = m.group(1).lower(), int(m.group(2))
        mo = _MONTH_ES.get(mname) or _MONTH_ES.get(mname[:3])
        if mo:
            return f'{y:04d}-{mo:02d}-01'

    return s


def parse_number_with_symbol(value: Any) -> tuple[str | None, str | None]:
    """Devuelve (valor_crudo, simbolo) siguiendo el estandar IIEG.

    Estandar: el valor se persiste sin separadores de miles, con punto como
    decimal. El simbolo `$` / `%` (prefijo o sufijo en la entrada) se extrae
    aparte. Si la entrada no es numerica (ej 'N/A', texto), se devuelve
    `(string_limpio, None)` para no perder data.
    """
    s = clean(value)
    if s is None:
        return (None, None)

    sym = None
    work = s

    m = _NUM_PREFIX_SYM.match(work)
    if m:
        sym = m.group(1)
        work = _NUM_PREFIX_SYM.sub('', work).strip()

    m = _NUM_SUFFIX_SYM.search(work)
    if m:
        sym = m.group(1)
        work = _NUM_SUFFIX_SYM.sub('', work).strip()

    work = work.replace(' ', '').replace(' ', '')
    if not work:
        return (s, sym)

    negative = work.startswith('-')
    if negative:
        work = work[1:]

    has_comma = ',' in work
    has_dot = '.' in work

    if has_comma and has_dot:
        if work.rfind('.') > work.rfind(','):
            cleaned = work.replace(',', '')
        else:
            cleaned = work.replace('.', '').replace(',', '.')
    elif has_comma:
        parts = work.split(',')
        cleaned = work.replace(',', '') if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and parts[0]) else work.replace(',', '.')
    elif has_dot:
        parts = work.split('.')
        cleaned = work.replace('.', '') if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and parts[0]) else work
    else:
        cleaned = work

    if not re.match(r'^\d+(\.\d+)?$', cleaned):
        return (s, None)

    if negative:
        cleaned = f'-{cleaned}'
    return (cleaned, sym)


def _is_numeralia_valor(field: str) -> bool:
    return field.startswith('numeralia_') and field.endswith('_valor')

TECHNICAL_FIELDS = {
    'layer_key',
    'workspace',
    'layer_name_db',
    'layer_name_usuario',
    'descripcion',
    'frecuencia',
    'fecha_ultima',
    'tipo_mapa',
    'tipo_mapa_enlace',
    'texto_leyenda',
    'tarjeta_punto_poligono',
    'link_final_capa',
    'downloadable',
    'fuentes_corto',
    'fuentes_largo',
    'fuentes_enlace',
    'metodologia_texto',
    'metodologia_archivo_enlace',
    'metadato_txt',
    'metadato_xlsx',
    'nombre_pie_numeralia',
    *(f'numeralia_0{i}_{kind}' for i in range(1, 9) for kind in ('nombre', 'valor', 'simbolo')),
}

MAPALAB_EXCEL_PRESET = {
    'Tema': '',
    'Subtema': '',
    'Link final de la capa': 'link_final_capa',
    'Capa Descargable': 'downloadable',
    'Capa finalizada en Geoserver': '',
    'Nombre de la capa en BD': 'layer_name_db',
    'Nombre de la capa en GeoServer': 'layer_key',
    'Nombre de la capa para el usuario': 'layer_name_usuario',
    'Columna en tabla a categorizar en mapa': '',
    'Descripción (número de caracteres recomendado 255-300)': 'descripcion',
    'Frecuencia de actualización': 'frecuencia',
    'Fecha de última actualización IIEG': 'fecha_ultima',
    'Numeralia 01 - Nombre': 'numeralia_01_nombre',
    'Numeralia 01 - Valor': 'numeralia_01_valor',
    'Numeralia 01 - Simbolo': 'numeralia_01_simbolo',
    'Numeralia 02 - Nombre': 'numeralia_02_nombre',
    'Numeralia 02 - Valor': 'numeralia_02_valor',
    'Numeralia 02 - Simbolo': 'numeralia_02_simbolo',
    'Numeralia 03 - Nombre': 'numeralia_03_nombre',
    'Numeralia 03 - Valor': 'numeralia_03_valor',
    'Numeralia 03 - Simbolo': 'numeralia_03_simbolo',
    'Numeralia 04 - Nombre': 'numeralia_04_nombre',
    'Numeralia 04 - Valor': 'numeralia_04_valor',
    'Numeralia 04 - Simbolo': 'numeralia_04_simbolo',
    'Numeralia 05 - Nombre': 'numeralia_05_nombre',
    'Numeralia 05 - Valor': 'numeralia_05_valor',
    'Numeralia 05 - Simbolo': 'numeralia_05_simbolo',
    'Numeralia 06 - Nombre': 'numeralia_06_nombre',
    'Numeralia 06 - Valor': 'numeralia_06_valor',
    'Numeralia 06 - Simbolo': 'numeralia_06_simbolo',
    'Numeralia 07 - Nombre': 'numeralia_07_nombre',
    'Numeralia 07 - Valor': 'numeralia_07_valor',
    'Numeralia 07 - Simbolo': 'numeralia_07_simbolo',
    'Numeralia 08 - Nombre': 'numeralia_08_nombre',
    'Numeralia 08 - Valor': 'numeralia_08_valor',
    'Numeralia 08 - Simbolo': 'numeralia_08_simbolo',
    'Nota al pie de numeralia (rango de fechas que comprende)': 'nombre_pie_numeralia',
    'Rangos de periodicidad (anual, mensual, a partir de que año-mes) BOTONES': '',
    'Metodología texto corto/ fórmula (número de caracteres)': 'metodologia_texto',
    'Metodología archivo/enlace': 'metodologia_archivo_enlace',
    'Fuentes (texto largo)': 'fuentes_largo',
    'Fuente (Texto corto Descargable PDF)': 'fuentes_corto',
    'Fuentes (enlace)': 'fuentes_enlace',
    'Tipo de mapa (IIEG/INEGI)': 'tipo_mapa',
    'Texto Leyenda Jurídico': 'texto_leyenda',
    'Tipo de mapa (enlace)': 'tipo_mapa_enlace',
    'Metadato .txt': 'metadato_txt',
    'Metadato .xlsx': 'metadato_xlsx',
    'Tarjeta de Geometría': 'tarjeta_punto_poligono',
    'Estatus': '',
    'Publicado Geoserver': '',
    'Publicado Mapalab': '',
}


class BulkIngestParseError(ValueError):
    pass


def clean(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() in ('nan', 'none', 'null'):
        return None
    return s


def parse_bool(value: Any) -> bool | None:
    if value is None or value == '':
        return None
    s = str(value).strip().lower()
    if s in TRUE_VALUES:
        return True
    if s in FALSE_VALUES:
        return False
    return None


def _build_fuentes(row: dict) -> dict | None:
    corto = clean(row.get('fuentes_corto'))
    largo = clean(row.get('fuentes_largo'))
    enlace = clean(row.get('fuentes_enlace'))
    if not any([corto, largo, enlace]):
        return None
    result: dict[str, str] = {}
    if corto:
        result['corto'] = corto
    if largo:
        result['largo'] = largo
    if enlace:
        result['enlace'] = enlace
    return result


def _build_metodologia(row: dict) -> dict | None:
    texto = clean(row.get('metodologia_texto'))
    enlace = clean(row.get('metodologia_archivo_enlace'))
    if not any([texto, enlace]):
        return None
    result: dict[str, str] = {}
    if texto:
        result['texto'] = texto
    if enlace:
        result['archivo_enlace'] = enlace
    return result


def _build_metadato(row: dict) -> list[dict] | None:
    items: list[dict] = []
    txt = clean(row.get('metadato_txt'))
    xlsx = clean(row.get('metadato_xlsx'))
    if txt:
        items.append({'nombre': 'Metadato TXT', 'enlace': txt})
    if xlsx:
        items.append({'nombre': 'Metadato XLSX', 'enlace': xlsx})
    return items or None


def _build_numeralia(row: dict) -> list[dict]:
    values: list[dict] = []
    for i in range(1, 9):
        valor = clean(row.get(f'numeralia_0{i}_valor'))
        nombre = clean(row.get(f'numeralia_0{i}_nombre'))
        simbolo = clean(row.get(f'numeralia_0{i}_simbolo'))
        if valor or nombre:
            values.append({
                'posicion': i,
                'valor': valor,
                'nombre': nombre,
                'simbolo': simbolo,
            })
    return values


def _read_csv(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('latin-1')
    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or [])]
    rows = [{(k or '').strip(): (v or '') for k, v in row.items()} for row in reader]
    return headers, rows


def _read_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise BulkIngestParseError(
            'openpyxl no instalado. Agregar al pyproject.toml.'
        ) from exc

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter: Iterator[tuple] = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(h).strip() if h is not None else '' for h in header_row]
    rows: list[dict] = []
    for raw in rows_iter:
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
            continue
        row: dict[str, str] = {}
        for i, value in enumerate(raw):
            if i >= len(headers) or not headers[i]:
                continue
            row[headers[i]] = '' if value is None else str(value).strip()
        rows.append(row)
    wb.close()
    return [h for h in headers if h], rows


def parse_source(filename: str, content: bytes) -> tuple[list[str], list[dict]]:
    lower = (filename or '').lower()
    if lower.endswith('.csv'):
        return _read_csv(content)
    if lower.endswith('.xlsx'):
        return _read_xlsx(content)
    raise BulkIngestParseError(
        f"Formato no soportado: '{filename}'. Sube CSV o XLSX."
    )


def apply_mapping(
    headers: list[str], rows: list[dict], mapping: dict[str, str]
) -> list[dict]:
    active = {src: dst for src, dst in mapping.items() if dst and src in headers}
    mapped_fields = set(active.values())
    normalized: list[dict] = []
    for raw in rows:
        mapped: dict[str, Any] = {}
        for src, dst in active.items():
            mapped[dst] = raw.get(src, '')

        row_out: dict[str, Any] = {}
        for field, value in mapped.items():
            if field == 'downloadable':
                parsed = parse_bool(value)
                if parsed is not None:
                    row_out[field] = parsed
            elif field == 'fecha_ultima':
                row_out[field] = parse_date_iso(value)
            elif _is_numeralia_valor(field):
                num, sym = parse_number_with_symbol(value)
                row_out[field] = num
                if sym:
                    sym_field = field.replace('_valor', '_simbolo')
                    if sym_field not in mapped_fields:
                        row_out[sym_field] = sym
            else:
                row_out[field] = clean(value)

        fuentes = _build_fuentes(row_out)
        metodologia = _build_metodologia(row_out)
        metadato = _build_metadato(row_out)
        numeralia = _build_numeralia(row_out)
        pie = row_out.get('nombre_pie_numeralia')

        out: dict[str, Any] = {}
        for k in (
            'layer_key',
            'workspace',
            'layer_name_db',
            'layer_name_usuario',
            'descripcion',
            'frecuencia',
            'fecha_ultima',
            'tipo_mapa',
            'tipo_mapa_enlace',
            'texto_leyenda',
            'tarjeta_punto_poligono',
            'link_final_capa',
        ):
            v = row_out.get(k)
            if v is not None:
                out[k] = v
        if 'downloadable' in row_out:
            out['downloadable'] = row_out['downloadable']
        if fuentes:
            out['fuentes'] = fuentes
        if metodologia:
            out['metodologia'] = metodologia
        if metadato:
            out['metadato'] = metadato
        if numeralia:
            out['values'] = numeralia
        if pie:
            out['pie_numeralia'] = pie

        key = out.get('layer_key')
        if key and not out.get('workspace') and ':' in key:
            out['workspace'] = key.split(':', 1)[0]

        normalized.append(out)
    return normalized


def unknown_targets(mapping: dict[str, str]) -> list[str]:
    return [
        dst for dst in mapping.values()
        if dst and dst not in TECHNICAL_FIELDS
    ]
