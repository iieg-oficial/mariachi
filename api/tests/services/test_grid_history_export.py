from datetime import datetime, timezone

from openpyxl import load_workbook

from app.services import grid_export
from app.services.grid_batch import (
    CellChange,
    GridField,
    GridSpec,
    GridTable,
    apply_cell_changes,
    diff_states,
    history_text,
    json_list_item_field,
)


class FakeResult:
    def __init__(self, rows, rowcount=1):
        self._rows = rows
        self.rowcount = rowcount

    def mappings(self):
        return self

    def first(self):
        return self._rows[0] if self._rows else None

    def all(self):
        return self._rows

    def scalar(self):
        if not self._rows:
            return None
        row = self._rows[0]
        return next(iter(row.values())) if isinstance(row, dict) else row


class FakeConnection:
    def __init__(self, rows_by_table=None, history_rows=None, history_exists='mapalab.grid_cell_history'):
        self.rows_by_table = rows_by_table or {}
        self.history_rows = history_rows or []
        self.history_exists = history_exists
        self.inserts = []
        self.statements = []

    def execute(self, statement, params=None):
        sql = str(statement)
        self.statements.append((sql, params))
        if 'to_regclass' in sql:
            return FakeResult([{'to_regclass': self.history_exists}])
        if 'grid_cell_history' in sql and sql.strip().startswith('INSERT'):
            self.inserts.extend(params if isinstance(params, list) else [params])
            return FakeResult([], rowcount=len(params) if isinstance(params, list) else 1)
        if 'grid_cell_history' in sql and sql.strip().startswith('SELECT'):
            return FakeResult(self.history_rows)
        if sql.strip().startswith('SELECT'):
            row = self.rows_by_table.get('metadata')
            return FakeResult([row] if row else [])
        return FakeResult([], rowcount=1)


METADATA_TABLE = GridTable(
    key='metadata',
    qualified_name='mapalab.layer_metadata',
    primary_key='layer_key',
    columns=('layer_key', 'descripcion', 'fuentes'),
    jsonb_columns=('fuentes',),
    audit_user_column='updated_by',
)


def _spec(history_table='mapalab.grid_cell_history'):
    fuentes_read, fuentes_write = json_list_item_field('fuentes', 0, 'corto')
    return GridSpec(
        key='layer-metadata',
        database='dataengine',
        row_key_field='layer_key',
        tables={'metadata': METADATA_TABLE},
        fields={
            'descripcion': GridField(
                key='descripcion',
                table='metadata',
                columns=('descripcion',),
                read=lambda row: row.get('descripcion'),
                write=lambda row, value: row.__setitem__('descripcion', value),
            ),
            'fuentes_corto': GridField(
                key='fuentes_corto',
                table='metadata',
                columns=('fuentes',),
                read=fuentes_read,
                write=fuentes_write,
            ),
        },
        fetch_rows=lambda conn, **kwargs: [],
        presence_scope='test',
        columns_meta=[
            {'key': 'descripcion', 'title': 'Descripción'},
            {'key': 'fuentes_corto', 'title': 'Fuente (corto)'},
        ],
        history_table=history_table,
    )


class TestHistoryText:
    def test_bool_legible(self):
        assert history_text(True) == 'true'

    def test_vacio_es_none(self):
        assert history_text('   ') is None

    def test_json_serializa(self):
        assert history_text({'a': 1}) == '{"a": 1}'


class TestRegistroDesdeElGrid:
    def test_registra_una_fila_por_celda(self):
        conn = FakeConnection({'metadata': {'layer_key': 'ws:capa', 'descripcion': 'vieja', 'fuentes': None}})
        apply_cell_changes(
            conn,
            _spec(),
            [CellChange('ws:capa', 'descripcion', 'vieja', 'nueva')],
            'ana@iieg.gob.mx',
        )
        assert len(conn.inserts) == 1
        entry = conn.inserts[0]
        assert entry['row_key'] == 'ws:capa'
        assert entry['column_key'] == 'descripcion'
        assert entry['from_value'] == 'vieja'
        assert entry['to_value'] == 'nueva'
        assert entry['changed_by'] == 'ana@iieg.gob.mx'
        assert entry['source'] == 'grid'

    def test_usa_la_clave_logica_no_el_jsonb(self):
        conn = FakeConnection({'metadata': {'layer_key': 'ws:capa', 'fuentes': [{'corto': 'CONAPO'}]}})
        apply_cell_changes(
            conn,
            _spec(),
            [CellChange('ws:capa', 'fuentes_corto', 'CONAPO', 'INEGI')],
            'ana@iieg.gob.mx',
        )
        entry = conn.inserts[0]
        assert entry['column_key'] == 'fuentes_corto'
        assert entry['from_value'] == 'CONAPO'
        assert entry['to_value'] == 'INEGI'

    def test_celda_en_conflicto_no_deja_historial(self):
        conn = FakeConnection({'metadata': {'layer_key': 'ws:capa', 'descripcion': 'la de otro'}})
        apply_cell_changes(
            conn,
            _spec(),
            [CellChange('ws:capa', 'descripcion', 'vieja', 'nueva')],
            'ana@iieg.gob.mx',
        )
        assert conn.inserts == []

    def test_sin_history_table_no_escribe(self):
        conn = FakeConnection({'metadata': {'layer_key': 'ws:capa', 'descripcion': 'vieja'}})
        apply_cell_changes(
            conn,
            _spec(history_table=None),
            [CellChange('ws:capa', 'descripcion', 'vieja', 'nueva')],
            'ana@iieg.gob.mx',
        )
        assert conn.inserts == []


class TestDiffStates:
    def test_detecta_cambios_por_campo_logico(self):
        spec = _spec()
        before = {'metadata': {'descripcion': 'vieja', 'fuentes': [{'corto': 'CONAPO'}]}}
        after = {'metadata': {'descripcion': 'nueva', 'fuentes': [{'corto': 'CONAPO'}]}}
        entries = diff_states(spec, 'ws:capa', before, after)
        assert len(entries) == 1
        assert entries[0]['column'] == 'descripcion'
        assert entries[0]['to_value'] == 'nueva'

    def test_sin_cambios_no_devuelve_nada(self):
        spec = _spec()
        state = {'metadata': {'descripcion': 'igual', 'fuentes': None}}
        assert diff_states(spec, 'ws:capa', state, state) == []

    def test_fila_nueva_registra_todos_los_campos_con_valor(self):
        spec = _spec()
        before = {'metadata': None}
        after = {'metadata': {'descripcion': 'nueva', 'fuentes': [{'corto': 'IIEG'}]}}
        entries = diff_states(spec, 'ws:capa', before, after)
        assert {entry['column'] for entry in entries} == {'descripcion', 'fuentes_corto'}


class TestExport:
    def test_hoja_de_datos_incluye_responsable(self):
        spec = _spec()
        rows = [{
            'layer_key': 'ws:capa',
            'descripcion': 'Una capa',
            'fuentes_corto': 'IIEG',
            'updated_by': 'ana@iieg.gob.mx',
            'updated_at': datetime(2026, 7, 28, 10, 30, tzinfo=timezone.utc),
        }]
        headers, body = grid_export.build_rows_sheet(spec, rows)
        assert headers[-2:] == ['Responsable', 'Última modificación']
        assert body[0][-2] == 'ana@iieg.gob.mx'
        assert body[0][-1] == '2026-07-28 10:30'

    def test_hoja_de_historial_traduce_campo_y_origen(self):
        spec = _spec()
        history = [{
            'row_key': 'ws:capa',
            'column_key': 'fuentes_corto',
            'from_value': 'CONAPO',
            'to_value': 'INEGI',
            'changed_by': 'ana@iieg.gob.mx',
            'changed_at': datetime(2026, 7, 28, 9, 0, tzinfo=timezone.utc),
            'source': 'grid',
        }]
        headers, body = grid_export.build_history_sheet(spec, history)
        assert headers[1] == 'Campo'
        assert body[0][1] == 'Fuente (corto)'
        assert body[0][6] == 'Tabla de captura'

    def test_xlsx_trae_las_dos_hojas(self):
        spec = _spec()
        rows_headers, rows_body = grid_export.build_rows_sheet(spec, [])
        history_headers, history_body = grid_export.build_history_sheet(spec, [])
        content = grid_export.to_xlsx([
            ('Metadatos', rows_headers, rows_body),
            ('Historial', history_headers, history_body),
        ])
        workbook = load_workbook(filename=__import__('io').BytesIO(content))
        assert workbook.sheetnames == ['Metadatos', 'Historial']

    def test_csv_lleva_bom_para_excel(self):
        content = grid_export.to_csv(['A', 'B'], [['1', '2']])
        assert content.startswith(b'\xef\xbb\xbf')
        assert b'1,2' in content

    def test_fetch_history_filtra_por_recurso(self):
        spec = _spec()
        conn = FakeConnection(history_rows=[{
            'row_key': 'ws:capa',
            'column_key': 'descripcion',
            'from_value': None,
            'to_value': 'x',
            'changed_by': 'ana',
            'changed_at': datetime(2026, 7, 28, tzinfo=timezone.utc),
            'source': 'grid',
        }])
        result = grid_export.fetch_history(conn, spec, date_from='2026-07-01', date_to='2026-07-31')
        assert len(result) == 1
        sql, params = conn.statements[-1]
        assert 'resource = :resource' in sql
        assert params['resource'] == 'layer-metadata'
        assert 'changed_at >= :date_from' in sql

    def test_sin_tabla_de_historial_exporta_igual(self):
        conn = FakeConnection(history_exists=None)
        assert grid_export.fetch_history(conn, _spec()) == []
